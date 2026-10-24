"""Turn an uploaded lecture file into slide images + extracted text.

PDF is the first-class path (PyMuPDF renders pages directly). PPTX extracts text
via python-pptx; for PPTX *images* you need LibreOffice on PATH to convert to PDF
first — see convert_pptx_to_pdf().
"""

from __future__ import annotations

import shutil
import subprocess
import uuid
from pathlib import Path

try:
    import pymupdf as fitz  # PyMuPDF >= 1.24.3
except ImportError:  # pragma: no cover
    import fitz  # older PyMuPDF; the `fitz` name is deprecated

from config import COURSES_DIR
from models import Lecture, Slide

RENDER_DPI = 150


def _lecture_dir(course_id: str, lecture_id: str) -> Path:
    d = COURSES_DIR / course_id / "lectures" / lecture_id / "slides"
    d.mkdir(parents=True, exist_ok=True)
    return d


def convert_pptx_to_pdf(pptx_path: Path, out_dir: Path) -> Path:
    """Shell out to LibreOffice. Returns the produced PDF path.

    Raises if soffice isn't installed — catch this and fall back to text-only.
    """
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice not found on PATH; cannot render PPTX slides.")
    out_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(pptx_path)],
        check=True,
        capture_output=True,
        timeout=180,
    )
    produced = out_dir / (pptx_path.stem + ".pdf")
    if not produced.exists():
        raise RuntimeError("LibreOffice ran but produced no PDF.")
    return produced


def _pptx_text_only(pptx_path: Path) -> list[str]:
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    out = []
    for slide in prs.slides:
        chunks = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                chunks.append(f"[speaker notes] {notes}")
        out.append("\n".join(c for c in chunks if c.strip()))
    return out


def ingest(
    source_path: Path,
    course_id: str,
    title: str = "",
    instructions: str = "",
) -> Lecture:
    """Read a lecture file, write slide PNGs, return a populated Lecture."""
    source_path = Path(source_path)
    lecture_id = uuid.uuid4().hex[:12]
    slides_dir = _lecture_dir(course_id, lecture_id)

    pdf_path = source_path
    text_fallback: list[str] = []

    if source_path.suffix.lower() in {".pptx", ".ppt"}:
        text_fallback = _pptx_text_only(source_path)
        try:
            pdf_path = convert_pptx_to_pdf(source_path, slides_dir.parent)
        except Exception:
            # No LibreOffice. Build a text-only lecture; the tutor still works,
            # the slide pane just shows text instead of the rendered slide.
            slides = [
                Slide(index=i, text=t, image_path="", width=0, height=0)
                for i, t in enumerate(text_fallback)
            ]
            return Lecture(
                id=lecture_id,
                title=title or source_path.stem,
                source_file=str(source_path),
                slides=slides,
                instructions=instructions,
            )

    doc = fitz.open(str(pdf_path))
    zoom = RENDER_DPI / 72.0
    matrix = fitz.Matrix(zoom, zoom)

    slides: list[Slide] = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=matrix)
        img_path = slides_dir / f"{i:03d}.png"
        pix.save(str(img_path))

        text = page.get_text("text").strip()
        if not text and i < len(text_fallback):
            text = text_fallback[i]

        slides.append(
            Slide(
                index=i,
                text=text,
                image_path=str(img_path),
                width=pix.width,
                height=pix.height,
            )
        )
    doc.close()

    return Lecture(
        id=lecture_id,
        title=title or source_path.stem,
        source_file=str(source_path),
        slides=slides,
        instructions=instructions,
    )
